import csv
import json
import os

from dotenv import load_dotenv
from openpyxl import load_workbook
import pytest

from utils.jsonhandling import jsonhandling


def test_jsonhandling():
    with open('testData\\creds.json') as data:
        formattedData = json.load(data)
        print(formattedData["positiveCreds"]["password"])


def test_jsonHandling2():
    data = jsonhandling('testData\\creds.json')
    print(data)


def test_csvhandling():
    with open('testData\\credentails.csv') as data:
        formattedData =csv.DictReader(data)
        values = []
        for i in formattedData:
            values.append(i)
        print(values[1]["username"])


#pip install openpyxl
# @pytest.mark.dh
def test_excelhandling():
    workbook = load_workbook("testData\\sample_creds.xlsx")
    sheet = workbook["Sheet2"]
    values = []
    for i in sheet[5]:
        print(i.value)
    # for i in sheet.iter_rows(min_row=2,values_only=True):
    #     values.append(i)
    # for i in sheet.iter_cols(min_col=1,values_only=True):
    #     values.append(i)
    # print(sheet["A4"].value)

    # print(values)

# @pytest.mark.dh
def test_excelhandling_write():
    workbook = load_workbook("testData\\sample_creds.xlsx")
    sheet = workbook["Sheet2"]
    # sheet["A4"]='tripura'
    # sheet.append(["new","new1"])
    sheet.delete_rows(4,1)
    workbook.save("testData\\sample_creds.xlsx")

@pytest.mark.skip
def test_copyPaste():
    workbook = load_workbook("testData\\sample_creds.xlsx")
    sheet = workbook["Sheet2"]
    sheet2 = workbook["sheet1"]
    startRow = sheet2.max_row+1
    for i in sheet.iter_rows(min_row=2, values_only=True):
        for cell1 in i:
            if "test" in cell1:
                sheet2.cell(row=startRow+cell1.row-1,cell=cell1.column).value= cell1
    workbook.save("testData\\sample_creds.xlsx")


def test_cli():
    usName = os.getenv("usname")
    pwValue = os.getenv("pwValues")
    print(usName)
    print(pwValue)

#pip install python-dotenv

def test_cli_env():
    load_dotenv(os.getenv("ENV_FILE"))
    usName = os.getenv("usname1")
    pwValue = os.getenv("pwValues1")
    url = os.getenv("url1")
    print(usName)
    print(pwValue)
    print(url)

@pytest.mark.parametrize("pwValues1", ["testpassword", "prodpassword", "devpassword"])
@pytest.mark.dh
def test_cli(pwValues1):
    usName = "test"
    pwValue = os.getenv(pwValues1)
    print(pwValue)








